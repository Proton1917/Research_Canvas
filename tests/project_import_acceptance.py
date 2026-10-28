import argparse
from contextlib import contextmanager
import io
import json
import mimetypes
import os
from pathlib import Path
import shutil
import subprocess
import time
import zipfile

from docx import Document
from markdown_it import MarkdownIt
from pptx import Presentation
from pptx.util import Inches
import requests
import toml
import xlsxwriter


ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / '.runtime/project-import'
RECORD = DATA / 'acceptance.json'
SOURCE = ROOT / 'docs/Research_Canvas_conversation_2026-10-25.pdf'


def save(record):
    RECORD.write_text(json.dumps(record, ensure_ascii=False, indent=2) + '\n')


class Client:
    def __init__(self, port):
        self.base = f'http://127.0.0.1:{port}/api'

    def request(self, method, path, expected=200, **kwargs):
        response = requests.request(method, self.base + path, timeout=180, **kwargs)
        assert response.status_code == expected, (path, response.status_code, response.text[:1800])
        return response.json() if response.headers.get('content-type', '').startswith('application/json') else response.content

    def get(self, path):
        return self.request('GET', path)

    def asset(self, id):
        response = requests.get(self.base + '/assets/' + id, timeout=60)
        assert response.status_code == 200
        return response.content

    def post(self, path, body, expected=200):
        return self.request('POST', path, expected, json=body)

    def upload(self, workspace, path):
        with path.open('rb') as stream:
            return self.request('POST', f'/workspaces/{workspace}/imports', 201, files={'file': (path.name, stream, mimetypes.guess_type(path)[0] or 'application/octet-stream')})

    def batch(self, project, files, directories=()):
        manifest = {'directories': list(directories), 'files': [{'field': str(i), 'path': path, **({'expected_version': version} if version is not None else {})} for i, (path, data, version) in enumerate(files)]}
        parts = [('manifest', (None, json.dumps(manifest, ensure_ascii=False), 'application/json'))]
        parts.extend((str(i), (Path(path).name, data, mimetypes.guess_type(path)[0] or 'application/octet-stream')) for i, (path, data, version) in enumerate(files))
        return self.request('POST', f'/projects/{project}/imports', files=parts)


@contextmanager
def service(name='main', port=8798):
    DATA.mkdir(parents=True, exist_ok=True)
    (DATA / 'empty-path').mkdir(exist_ok=True)
    config = toml.load(ROOT / 'config/server.example.toml')
    config.update(port=port, database=str(DATA / f'{name}.sqlite3'), project_page_size=7)
    path = DATA / f'{name}.toml'
    path.write_text(toml.dumps(config))
    with (DATA / f'{name}.log').open('a') as log:
        process = subprocess.Popen([str(ROOT / '.runtime/bin/research-canvas')], cwd=ROOT, env={**os.environ, 'RC_CONFIG': str(path), 'PATH': str(DATA / 'empty-path'), 'TMPDIR': str(DATA)}, stdout=log, stderr=log)
        try:
            client = Client(port)
            for _ in range(100):
                assert process.poll() is None, f'{name}.log'
                try:
                    if requests.get(client.base + '/health', timeout=1).ok:
                        break
                except requests.ConnectionError:
                    pass
                time.sleep(0.1)
            else:
                raise AssertionError('服务启动超时')
            yield client
        finally:
            process.terminate()
            process.wait(timeout=30)


def fixtures(record):
    folder = DATA / 'inputs'
    folder.mkdir(parents=True, exist_ok=True)
    readme = (ROOT / 'README.md').read_text()
    paragraph = next(line for line in readme.splitlines() if line.startswith('Research Canvas 是'))
    files = sorted(path for path in (ROOT / 'docs').rglob('*') if path.is_file())
    inventory = [(path.relative_to(ROOT).as_posix(), path.stat().st_size) for path in files]
    (folder / 'project.note').write_text(readme)
    (folder / 'project.json').write_bytes((ROOT / 'package.json').read_bytes())
    (folder / 'README.txt').write_bytes(readme.encode('utf-16'))
    notebook = {'nbformat': 4, 'nbformat_minor': 5, 'metadata': {'language_info': {'name': 'python'}}, 'cells': [
        {'cell_type': 'markdown', 'metadata': {}, 'source': readme},
        {'cell_type': 'code', 'metadata': {}, 'source': (ROOT / 'scripts/release_targets.py').read_text(), 'outputs': [], 'execution_count': None},
    ]}
    (folder / 'research.ipynb').write_text(json.dumps(notebook, ensure_ascii=False))
    document = Document()
    document.add_heading('Research Canvas', level=1)
    document.add_paragraph(paragraph)
    document.add_paragraph((ROOT / 'docs/backend-integration.md').read_text().splitlines()[0])
    document.save(folder / 'research.docx')
    with xlsxwriter.Workbook(folder / 'project-files.xlsx') as workbook:
        sheet = workbook.add_worksheet('项目文件')
        sheet.write_row(0, 0, ['文件', '字节数'])
        for row, values in enumerate(inventory, 1):
            sheet.write_row(row, 0, values)
        total = sum(size for _, size in inventory)
        sheet.write(len(inventory) + 1, 0, '合计')
        formula = f'SUM(B2:B{len(inventory) + 1})'
        sheet.write_formula(len(inventory) + 1, 1, '=' + formula, None, total)
        version = workbook.add_worksheet('应用版本')
        version.write_row(0, 0, ['名称', '版本'])
        version.write_row(1, 0, ['research-canvas', json.loads((ROOT / 'package.json').read_text())['version']])
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = 'Research Canvas'
    slide.placeholders[1].text = paragraph
    slide.notes_slide.notes_text_frame.text = readme.splitlines()[4]
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = '项目文件'
    table = slide.shapes.add_table(4, 2, Inches(0.5), Inches(1.5), Inches(9), Inches(3)).table
    for row, values in enumerate([('文件', '字节数'), *inventory[:3]]):
        for column, value in enumerate(values):
            table.cell(row, column).text = str(value)
    slide.notes_slide.notes_text_frame.text = '来源：' + ', '.join(name for name, _ in inventory[:3])
    presentation.save(folder / 'research.pptx')
    soffice = shutil.which('soffice')
    assert soffice
    profile = (DATA / 'office-profile').as_uri()
    environment = {**os.environ, 'TMPDIR': str(DATA), 'XDG_CACHE_HOME': str(DATA / 'cache')}
    for source, format_name in [('project-files.xlsx', 'xls:MS Excel 97'), ('research.pptx', 'ppt:MS PowerPoint 97')]:
        subprocess.run([soffice, '-env:UserInstallation=' + profile, '--headless', '--convert-to', format_name, '--outdir', str(folder), str(folder / source)], env=environment, check=True, timeout=60)
        assert (folder / (Path(source).stem + '.' + format_name.split(':')[0])).is_file()
    font = next((ROOT / 'dist/assets').glob('*.woff2'))
    shutil.copyfile(font, folder / font.name)
    record.update(fixtures='passed', source_paragraph=paragraph, inventory_total=total, inventory_formula=formula, table_source=inventory[0][0], attachment=font.name)
    save(record)


def formats(client, record):
    if 'workspace_id' not in record:
        record['workspace_id'] = client.post('/workspaces', {'name': '文件夹与 Office 导入验收'}, 201)['id']
        record['documents'] = {}
        save(record)
    workspace = record['workspace_id']
    saved = {document['title']: document for document in client.get('/workspaces/' + workspace)['documents']}
    capabilities = client.get('/imports/capabilities')
    assert capabilities['pdf_text'] == 'builtin'
    for path in sorted((DATA / 'inputs').iterdir()):
        if path.name in record['documents']:
            continue
        if path.name in saved:
            doc = saved[path.name]
            result = {'document': doc, 'asset': {'id': doc['metadata']['source_asset_id']}, 'text_status': doc['metadata']['text_status'], 'message': doc['metadata']['import_message']}
        else:
            result = client.upload(workspace, path)
        doc = result['document']
        body = '\n'.join(block['text'] for block in doc['content']['blocks'])
        if path.name == record['attachment']:
            assert result['text_status'] == 'attachment'
        else:
            assert result['text_status'] in {'ready', 'partial'}, result
        if path.suffix in {'.xlsx', '.xls'}:
            assert '项目文件' in body and '应用版本' in body and str(record['inventory_total']) in body and record['inventory_formula'] in body, body[-1800:]
            assert len(doc['metadata']['import_details']['sheets']) == 2
        if path.suffix in {'.pptx', '.ppt'}:
            assert '幻灯片 1' in body and 'Research Canvas' in body and '来源：' in body, body
            if path.suffix == '.pptx':
                tokens = MarkdownIt('commonmark').enable('table').parse(body)
                text = '\n'.join(child.content for token in tokens for child in (token.children or []) if child.type == 'text')
                assert record['table_source'] in text and '| ' in body, body
            assert len(doc['metadata']['import_details']['slides']) == 2
        if path.suffix in {'.docx', '.note', '.txt', '.ipynb'}:
            assert record['source_paragraph'] in body
        if path.suffix == '.json':
            assert body == (ROOT / 'package.json').read_text() and doc['content']['markdown'].startswith('```json')
        assert client.asset(result['asset']['id']) == path.read_bytes()
        record['documents'][path.name] = doc['id']
        save(record)
    pdf = client.upload(workspace, SOURCE)
    assert pdf['text_status'] == 'ready', pdf['message']
    assert len(pdf['document']['content']['blocks']) == 7
    assert [block['page'] for block in pdf['document']['content']['blocks']] == list(range(1, 8))
    assert 'Research Canvas' in '\n'.join(block['text'] for block in pdf['document']['content']['blocks'])
    record['pdf_document'] = pdf['document']['id']
    record['formats'] = 'passed'
    save(record)


def tree_entries(client, project):
    pending = ['']
    entries = []
    pages = 0
    while pending:
        parent = pending.pop()
        after = ''
        while True:
            result = client.request('GET', f'/projects/{project}/entries', params={'parent': parent, 'after': after})
            pages += 1
            entries.extend(result['entries'])
            pending.extend(entry['path'] for entry in result['entries'] if entry['kind'] == 'directory')
            after = result['next_cursor']
            if after is None:
                break
    return entries, pages


def projects(client, record):
    workspace = record['workspace_id']
    record['projects'] = {}
    for name, source in [('研究文档', ROOT / 'docs'), ('模型提示词', ROOT / 'config/prompts')]:
        project = client.post(f'/workspaces/{workspace}/projects', {'name': name}, 201)
        files = [(path.relative_to(source).as_posix(), path.read_bytes(), None) for path in sorted(source.rglob('*')) if path.is_file()]
        directories = [path.relative_to(source).as_posix() for path in source.rglob('*') if path.is_dir()]
        batch = []
        byte_count = 0
        results = []
        started = time.monotonic()
        for item in files:
            if batch and (byte_count + len(item[1]) > 16 * 1024 * 1024 or len(batch) >= 24):
                results.extend(client.batch(project['id'], batch)['results'])
                batch, byte_count = [], 0
            batch.append(item)
            byte_count += len(item[1])
        results.extend(client.batch(project['id'], batch, directories)['results'])
        elapsed = time.monotonic() - started
        assert len(results) == len(files) and all(item['status'] == 'imported' for item in results)
        assert all(item['entry']['text_status'] != 'error' for item in results), [(item['path'], item['entry']['message']) for item in results if item['entry']['text_status'] == 'error']
        entries, pages = tree_entries(client, project['id'])
        assert {entry['path'] for entry in entries if entry['kind'] == 'file'} == {path for path, _, _ in files}
        assert {entry['path'] for entry in entries if entry['kind'] == 'directory'} == set(directories)
        for (path, data, _), imported in zip(files, results):
            assert imported['entry']['path'] == path
            assert client.asset(imported['entry']['asset_id']) == data
        record['projects'][name] = {'id': project['id'], 'count': len(files), 'bytes': sum(len(data) for _, data, _ in files), 'seconds': elapsed, 'pages': pages}
        save(record)
    project_ids = {project['id'] for project in record['projects'].values()}
    summaries = client.get(f'/workspaces/{workspace}/projects')
    assert {item['project']['id'] for item in summaries} == project_ids
    lightweight = client.request('GET', '/workspaces/' + workspace, params={'project_documents': 'false'})
    assert len(lightweight['documents']) == len(record['documents']) + 1
    record['project_tree'] = 'passed'
    save(record)


def reimport(client, record):
    project = record['projects']['研究文档']['id']
    source = ROOT / 'docs/backend.md'
    current = client.batch(project, [('backend.md', source.read_bytes(), None)])['results'][0]
    assert current['status'] == 'unchanged'
    document_id = current['entry']['document_id']
    historical = subprocess.check_output(['git', 'show', 'v0.1.0:docs/backend.md'], cwd=ROOT)
    assert historical != source.read_bytes()
    updated = client.batch(project, [('backend.md', historical, None)])['results'][0]
    assert updated['status'] == 'updated' and updated['entry']['document_id'] == document_id and updated['entry']['imported_version'] == 2
    edited = client.request('PATCH', '/documents/' + document_id, json={'expected_version': 2, 'title': '研究文档/backend.md（研究批注）'})
    results = client.batch(project, [('backend.md', source.read_bytes(), None), ('reference/README.md', (ROOT / 'README.md').read_bytes(), None)], ['reference/empty'])['results']
    assert results[0]['status'] == 'conflict' and results[1]['status'] == 'imported'
    assert client.get('/documents/' + document_id)['title'] == edited['title']
    accepted = client.batch(project, [('backend.md', source.read_bytes(), edited['current_version'])])['results'][0]
    assert accepted['status'] == 'updated' and accepted['entry']['imported_version'] == 4
    context = client.post(f"/workspaces/{record['workspace_id']}/contexts", {'question': '读取导入文稿的后端配置说明。', 'document_ids': [document_id]}, 201)
    assert source.read_text() in context['sources'][0]['content']
    assert len(client.get('/documents/' + document_id)['metadata']['source_versions']) == 3
    record.update(reimport='passed', updated_document=document_id)
    save(record)


def archive(client, record):
    workspace = record['workspace_id']
    exported = client.get(f'/workspaces/{workspace}/export')
    assert exported['format_version'] == 4 and len(exported['projects']) == 2
    (DATA / 'workspace.json').write_text(json.dumps(exported, ensure_ascii=False))
    data = client.get(f'/workspaces/{workspace}/export.zip')
    with zipfile.ZipFile(io.BytesIO(data)) as zipped:
        assert json.loads(zipped.read('workspace.json')) == exported
        project = record['projects']['研究文档']['id']
        assert zipped.read(f'projects/{project}/backend.md') == (ROOT / 'docs/backend.md').read_bytes()
        assert zipped.getinfo(f'projects/{project}/reference/empty/').is_dir()
    with service('restored', 8799) as restored:
        restored.post('/workspaces/import', exported, 201)
        assert restored.get(f'/workspaces/{workspace}/export') == exported
        entries, _ = tree_entries(restored, record['projects']['研究文档']['id'])
        assert any(entry['path'] == 'reference/empty' for entry in entries)
    record.update(archive='passed', archive_bytes=(DATA / 'workspace.json').stat().st_size)
    save(record)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('phase', choices=['fixtures', 'formats', 'projects', 'reimport', 'archive'])
    phase = parser.parse_args().phase
    DATA.mkdir(parents=True, exist_ok=True)
    record = json.loads(RECORD.read_text()) if RECORD.exists() else {}
    key = 'project_tree' if phase == 'projects' else phase
    if record.get(key) == 'passed':
        raise SystemExit(phase + ' 已通过，复用记录。')
    if phase == 'fixtures':
        fixtures(record)
    else:
        with service() as client:
            globals()[phase](client, record)
    print(phase + ': passed', flush=True)


if __name__ == '__main__':
    main()
